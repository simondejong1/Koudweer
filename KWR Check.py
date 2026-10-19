import os
import requests
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import matplotlib.pyplot as plt

# 1. Definieer de G4 steden met hun coördinaten
g4_steden = {
    "Amsterdam": {"latitude": 52.3676, "longitude": 4.9041},
    "Rotterdam": {"latitude": 51.9244, "longitude": 4.4777},
    "Den Haag": {"latitude": 52.0705, "longitude": 4.3007},
    "Utrecht": {"latitude": 52.0907, "longitude": 5.1214}
}

# API-sjabloon voor een 15-daagse uurlijkse verwachting van de gevoelstemperatuur
url_template = "https://api.open-meteo.com/v1/forecast?latitude={lat}&longitude={lon}&hourly=apparent_temperature&forecast_days=15"

all_data = []

print("Bezig met ophalen van de weersverwachtingen...")

# 2. Loop door elke stad en haal de data op
for stad, coords in g4_steden.items():
    url = url_template.format(lat=coords["latitude"], lon=coords["longitude"])
    try:
        response = requests.get(url)
        if response.status_code == 200:
            data = response.json()
            hourly = data.get("hourly", {})
            times = hourly.get("time", [])
            apparent_temps = hourly.get("apparent_temperature", [])
            
            # Koppel tijdstip en gevoelstemperatuur aan elkaar
            for t, temp in zip(times, apparent_temps):
                all_data.append({
                    "Stad": stad,
                    "Datum_Tijd": t,
                    "Gevoelstemperatuur_Celsius": temp
                })
            print(f" [OK] {stad} succesvol opgehaald.")
        else:
            print(f" [FOUT] Kon geen data ophalen voor {stad} (Status: {response.status_code})")
    except Exception as e:
        print(f" [FOUT] Er ging iets mis bij {stad}: {e}")

print(f"Aantal verzamelde datapunten: {len(all_data)}")

# 3. Verwerk de data en sla op als Excel-bestand (.xlsx) op het Bureaublad
if all_data:
    df = pd.DataFrame(all_data)
    
    # Bepaal het pad naar het Bureaublad
    bureaublad = os.path.join(os.path.expanduser("~"), "Desktop")
    excel_naam = os.path.join(bureaublad, "g4_gevoelstemperatuur_15dagen.xlsx")
    
    # Exporteer naar Excel
    df.to_excel(excel_naam, index=False, sheet_name="Gevoelstemperatuur")
    
    # 4. Pas kleuring toe via openpyxl (Rood voor < 0, Oranje voor 0 t/m 5)
    wb = load_workbook(excel_naam)
    ws = wb.active
    
    rood_fill = PatternFill(start_color="FF4C4C", end_color="FF4C4C", fill_type="solid")
    oranje_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    
    for row in range(2, ws.max_row + 1):
        cel = ws.cell(row=row, column=3)
        cel_waarde = cel.value
        if cel_waarde is not None:
            if cel_waarde < 0:
                cel.fill = rood_fill
            elif 0 <= cel_waarde <= 5:
                cel.fill = oranje_fill
                
    wb.save(excel_naam)
    print(f"\nExcel-bestand met rood/oranje markeringen opgeslagen op Bureaublad.")

    # 5. Automatisch een grafiek plotten
    print("Bezig met genereren van de grafiek...")
    
    df['Datum_Tijd'] = pd.to_datetime(df['Datum_Tijd'])

    plt.figure(figsize=(12, 6))
    
    for stad, group in df.groupby('Stad'):
        plt.plot(group['Datum_Tijd'], group['Gevoelstemperatuur_Celsius'], label=stad, linewidth=1.5)

    plt.title('15-daagse Gevoelstemperatuur Verwachting G4', fontsize=14, fontweight='bold')
    plt.xlabel('Datum en Tijd', fontsize=12)
    plt.ylabel('Gevoelstemperatuur (°C)', fontsize=12)
    plt.axhline(0, color='blue', linestyle='--', linewidth=0.8, alpha=0.5, label='0°C')
    plt.axhline(5, color='orange', linestyle='--', linewidth=0.8, alpha=0.5, label='5°C')
    
    plt.legend(loc='upper right')
    plt.grid(True, linestyle=':', alpha=0.6)
    plt.tight_layout()

    grafiek_naam = os.path.join(bureaublad, "g4_gevoelstemperatuur_grafiek.png")
    plt.savefig(grafiek_naam, dpi=300)
    print(f"Grafiek als afbeelding opgeslagen op Bureaublad: {grafiek_naam}")

    plt.show()

else:
    print("Geen data verzameld.")

input("\nDruk op Enter om dit venster te sluiten...")
